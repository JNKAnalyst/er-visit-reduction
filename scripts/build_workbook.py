"""Generate a starter ER_Visit_Reduction_Model workbook from CSV templates.

Produces an .xlsx skeleton with the Config, Baseline, Dashboard, and Sensitivity
sheets pre-populated. The VBA modules under ``model/vba/`` must be imported
manually in the Excel VBA editor (Alt+F11 -> File -> Import File...) and the
file then saved as ``.xlsm``. This split keeps the VBA source readable in git
and avoids platform-specific binary regeneration.

Usage:
    python scripts/build_workbook.py
"""
from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
TEMPLATES = ROOT / "templates"
OUT = ROOT / "model" / "ER_Visit_Reduction_Model.xlsx"

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")


def _read_csv(path: Path) -> list[list[str]]:
    with path.open(newline="") as f:
        return list(csv.reader(f))


def _style_header(ws, row: int, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL


def _autosize(ws) -> None:
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(width + 2, 40)


def build() -> Path:
    wb = Workbook()

    # Config
    cfg = wb.active
    cfg.title = "Config"
    for row in _read_csv(TEMPLATES / "config_template.csv"):
        cfg.append(row)
    _style_header(cfg, 1, 3)
    _autosize(cfg)

    # Baseline
    base = wb.create_sheet("Baseline")
    for row in _read_csv(TEMPLATES / "baseline_visits_template.csv"):
        base.append(row)
    _style_header(base, 1, 4)
    _autosize(base)

    # Dashboard
    dash = wb.create_sheet("Dashboard")
    dash["A1"] = "ER Visit Reduction — Dashboard"
    dash["A1"].font = Font(bold=True, size=14)
    dash["A3"] = "Click 'Run Analysis' (assigned to modDashboard.RunAnalysis) after importing VBA modules."
    dash["A10"] = "Scenario"
    dash["B10"] = "Reduction"
    dash["C10"] = "Visits Avoided"
    dash["D10"] = "Monthly Savings"
    dash["E10"] = "Annual Savings"
    _style_header(dash, 10, 5)
    _autosize(dash)

    # Sensitivity
    sens = wb.create_sheet("Sensitivity")
    sens["A1"] = "Parameter"
    sens["B1"] = "Delta"
    sens["C1"] = "Combined Reduction"
    sens["D1"] = "Annual Savings"
    _style_header(sens, 1, 4)
    _autosize(sens)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    return OUT


if __name__ == "__main__":
    out = build()
    print(f"Wrote {out.relative_to(ROOT)}")
